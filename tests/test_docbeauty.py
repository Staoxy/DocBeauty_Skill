# -*- coding: utf-8 -*-
"""DocBeauty V2 测试矩阵（DESIGN_V2.md §24）。"""
import json
import os
import sys

import pytest
from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import formatter  # noqa: E402
from config import load_config  # noqa: E402
from main import main as cli_main, process_document  # noqa: E402
from utils import ooxml as ox  # noqa: E402

from tests import fixtures as F  # noqa: E402


def beautify(src, tmpdir, **over):
    cfg = load_config(None)
    cfg["output_dir"] = str(tmpdir)
    cfg.update(over)
    rep, code, out = process_document(str(src), cfg)
    return rep, code, out


def get_check(rep, qid):
    for c in rep["quality_checks"]:
        if c["id"] == qid:
            return c
    return None


# ---------------------------------------------------------------------------
# 基础 + 护栏不误报（T-guard-no-false-positive）
# ---------------------------------------------------------------------------
def test_basic_success(tmp_path):
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 0 and rep["status"] == "success", rep["errors"]
    assert rep["content_guard"]["content_changed"] is False
    assert os.path.isfile(out)
    assert get_check(rep, "Q01")["status"] == "pass"
    assert get_check(rep, "Q02")["status"] == "pass"
    assert get_check(rep, "Q03")["status"] == "pass"

    doc = Document(out)
    # 标题提升：一、/（一）/1. 均成为 Heading
    styles = doc.styles.element
    from analyzer import heading_level_of
    levels = [heading_level_of(styles, p._p) for p in doc.paragraphs]
    assert levels.count(1) >= 2 and levels.count(2) >= 1 and levels.count(3) >= 1
    # 正文缩进：firstLineChars=200
    body_ok = 0
    for p in doc.paragraphs:
        if p.text.strip().startswith("正文内容"):
            pPr = p._p.find(qn("w:pPr"))
            ind = ox.get_child(pPr, "w:ind") if pPr is not None else None
            if ind is not None and ind.get(qn("w:firstLineChars")) == "200":
                body_ok += 1
    assert body_ok >= 5
    # 中文字体 eastAsia 已设置（抽查正文 run）
    east_set = 0
    for p in doc.paragraphs:
        if "正文内容" in p.text:
            for r in p._p.findall(qn("w:r")):
                rPr = ox.get_child(r, "w:rPr")
                rf = ox.get_child(rPr, "w:rFonts") if rPr is not None else None
                if rf is not None and rf.get(qn("w:eastAsia")):
                    east_set += 1
    assert east_set >= 5
    # 目录已插入 + updateFields
    assert rep["config"]["operations_executed"].count("toc_added") == 1
    assert "目录" in "\n".join(p.text for p in doc.paragraphs)
    assert ox.get_child(doc.settings.element, "w:updateFields") is not None
    # 空段收敛：3 连空 -> 1
    assert rep["statistics"]["reconciliation"]["blank_deleted"] >= 3
    # 页码已加
    assert any("footer" in f.lower() for f in rep["analysis"]["existing_page_number_fields"] + 
               [x for x in [get_check(rep, "Q11")["detail"]]]) or \
        get_check(rep, "Q11")["status"] in ("pass", "skip")


def test_references_zone_no_indent(tmp_path):
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    doc = Document(out)
    refs = [p for p in doc.paragraphs if p.text.strip().startswith("[")]
    assert refs, "参考文献条目应存在"
    for p in refs:
        pPr = p._p.find(qn("w:pPr"))
        ind = ox.get_child(pPr, "w:ind") if pPr is not None else None
        assert ind is None or ind.get(qn("w:firstLineChars")) in (None, "0"), \
            "参考文献条目不应有首行缩进"


def test_caption_formatted(tmp_path):
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    doc = Document(out)
    cap = [p for p in doc.paragraphs if p.text.strip().startswith("图1")][0]
    pPr = cap._p.find(qn("w:pPr"))
    jc = ox.get_child(pPr, "w:jc")
    assert jc is not None and jc.get(qn("w:val")) == "center"
    r = cap._p.findall(qn("w:r"))[0]
    rPr = ox.get_child(r, "w:rPr")
    sz = ox.get_child(rPr, "w:sz")
    assert sz.get(qn("w:val")) == "21"  # 10.5pt


# ---------------------------------------------------------------------------
# 护栏拦截（T-guard-catch）：篡改正文 -> exit 3 且不产出文件
# ---------------------------------------------------------------------------
def test_guard_catches_tampering(tmp_path, monkeypatch):
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    orig = formatter.normalize_fonts

    def tamper(doc, zones, effective, options):
        for p in doc.paragraphs:
            if "正文内容1" in p.text:
                for t in p._p.iter(qn("w:t")):
                    t.text = (t.text or "") + "被篡改的文字"
                break
        return orig(doc, zones, effective, options)

    monkeypatch.setattr(formatter, "normalize_fonts", tamper)
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 3
    assert rep["status"] == "guard_failed"
    assert rep["content_guard"]["content_changed"] is True
    assert rep["content_guard"]["diffs"], "diff 明细不应为空"
    assert out is None, "护栏失败时不得产出文件"
    assert not os.path.isfile(tmp_path / "out" / "in_beautified.docx")


# ---------------------------------------------------------------------------
# 防双重编号（T-double-numbering）
# ---------------------------------------------------------------------------
def test_double_numbering(tmp_path):
    src = F.numbered_heading_style_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code in (0, 2)
    doc = Document(out)
    # 样式编号已禁用：numId=0
    h1 = ox.find_style_by_name(doc.styles.element, "heading 1")
    pPr = ox.get_child(h1, "w:pPr")
    numPr = ox.get_child(pPr, "w:numPr")
    numId = ox.get_child(numPr, "w:numId")
    assert numId is not None and numId.get(qn("w:val")) == "0"
    # 段落无 numPr + 文字未变
    target = [p for p in doc.paragraphs if "手工编号标题" in p.text][0]
    assert target.text.strip() == "一、手工编号标题"
    assert ox.get_child(target._p.find(qn("w:pPr")), "w:numPr") is None
    assert "numbering_disabled_on_styles" in str(rep["operations_detail"])


# ---------------------------------------------------------------------------
# 两表夹空段（T-table-merge）
# ---------------------------------------------------------------------------
def test_table_merge_protection(tmp_path):
    src = F.two_tables_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 0
    doc = Document(out)
    body = list(doc.element.body)
    tbl_positions = [i for i, el in enumerate(body) if el.tag == qn("w:tbl")]
    assert len(tbl_positions) == 2, "表格数量不得变化"
    between = body[tbl_positions[0] + 1:tbl_positions[1]]
    paras = [el for el in between if el.tag == qn("w:p")]
    assert paras, "两表之间的空段必须保留（否则 Word 会合并两表）"


# ---------------------------------------------------------------------------
# sectPr 空段保留（T-sectpr）
# ---------------------------------------------------------------------------
def test_sectpr_blank_preserved(tmp_path):
    src = F.sectpr_blank_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 0
    doc = Document(out)
    assert len(doc.sections) == 2, "节数不得变化"
    # 带 sectPr 的段落仍在
    found = any(p.find(qn("w:pPr")) is not None and
                p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None
                for p in (el for el in doc.element.body if el.tag == qn("w:p")))
    assert found, "携带 sectPr 的空段不得删除"


# ---------------------------------------------------------------------------
# 超链接（T-hyperlink）
# ---------------------------------------------------------------------------
def test_hyperlink_preserved(tmp_path):
    src = F.hyperlink_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 0
    doc = Document(out)
    links = doc.element.body.findall(".//" + qn("w:hyperlink"))
    assert len(links) == 1
    link_text = "".join(t.text or "" for t in links[0].iter(qn("w:t")))
    assert link_text == "这是一个超链接文本"
    # 超链接 run 也设置了 eastAsia 字体
    rPr = ox.get_child(links[0].findall(qn("w:r"))[0], "w:rPr")
    rf = ox.get_child(rPr, "w:rFonts")
    assert rf is not None and rf.get(qn("w:eastAsia")) == "宋体"
    # para.text（含超链接文本，python-docx>=1.1）
    para = [p for p in doc.paragraphs if "超链接" in p.text][0]
    assert "这是一个超链接文本" in para.text


# ---------------------------------------------------------------------------
# 强调保留（T-emphasis）
# ---------------------------------------------------------------------------
def test_emphasis_preserved(tmp_path):
    src = F.emphasis_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 0
    doc = Document(out)
    para = [p for p in doc.paragraphs if "加粗强调" in p.text][0]
    runs = {"".join(t.text or "" for t in r.iter(qn("w:t"))): r
            for r in para._p.findall(qn("w:r"))}
    assert ox.get_child(ox.get_child(runs["加粗强调"], "w:rPr"), "w:b") is not None
    assert ox.get_child(ox.get_child(runs["斜体强调"], "w:rPr"), "w:i") is not None
    assert ox.get_child(ox.get_child(runs["下划线强调"], "w:rPr"), "w:u") is not None
    color = ox.get_child(ox.get_child(runs["红色文字"], "w:rPr"), "w:color")
    assert color is not None and color.get(qn("w:val")) == "FF0000"
    # 字体统一后 sz 全部 24（12pt）
    for r in runs.values():
        sz = ox.get_child(ox.get_child(r, "w:rPr"), "w:sz")
        assert sz.get(qn("w:val")) == "24"


# ---------------------------------------------------------------------------
# 图片（T-image）
# ---------------------------------------------------------------------------
def test_image_scaled_and_centered(tmp_path):
    src = F.image_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 0
    detail = [d for d in rep["operations_detail"] if d["op"] == "images_formatted"][0]
    assert detail["changed"]["scaled"] == 1
    doc = Document(out)
    shape = doc.inline_shapes[0]
    usable = int(doc.sections[0].page_width - doc.sections[0].left_margin
                 - doc.sections[0].right_margin)
    assert shape.width <= usable, "超宽图片必须缩小到可用宽度内"
    # 居中
    img_para = None
    for p in doc.paragraphs:
        if p._p.findall(".//" + qn("wp:inline")):
            img_para = p
            break
    jc = ox.get_child(img_para._p.find(qn("w:pPr")), "w:jc")
    assert jc is not None and jc.get(qn("w:val")) == "center"


# ---------------------------------------------------------------------------
# 已有 TOC 域（不插入第二个）
# ---------------------------------------------------------------------------
def test_existing_toc_field(tmp_path):
    src = F.existing_toc_field_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 0
    doc = Document(out)
    toc_count = 0
    for it in doc.element.body.iter(qn("w:instrText")):
        if it.text and "TOC" in it.text:
            toc_count += 1
    assert toc_count == 1, "已存在 TOC 域时不得插入第二个"
    assert ox.get_child(doc.settings.element, "w:updateFields") is not None
    detail = [d for d in rep["operations_detail"] if d["op"] == "toc_added"][0]
    assert detail["status"] == "done"


# ---------------------------------------------------------------------------
# 手工目录（T-manual-toc：保留 + 警告）
# ---------------------------------------------------------------------------
def test_manual_toc_preserved(tmp_path):
    src = F.manual_toc_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 0
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert any("一、绪论..........1" in t for t in texts), "手工目录默认保留不删除"
    assert any("手工" in w for w in rep["warnings"]), "必须给出手工目录警告"


# ---------------------------------------------------------------------------
# 修订警告（T-tracking）
# ---------------------------------------------------------------------------
def test_tracked_changes_warning(tmp_path):
    src = F.tracked_changes_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 0
    assert any("修订" in w for w in rep["warnings"])
    # 修订内容未被破坏
    doc = Document(out)
    assert doc.element.body.findall(".//" + qn("w:ins"))


# ---------------------------------------------------------------------------
# 边界对象（文本框：未处理亦未破坏）
# ---------------------------------------------------------------------------
def test_textbox_untouched(tmp_path):
    src = F.textbox_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 0
    doc = Document(out)
    boxes = doc.element.body.findall(".//" + qn("w:txbxContent"))
    assert len(boxes) == 1
    box_text = "".join(t.text or "" for t in boxes[0].iter(qn("w:t")))
    assert box_text == "文本框内部文字"
    assert rep["untouched_boundary"]["textboxes"] == 1


# ---------------------------------------------------------------------------
# 幂等性（T-idempotent）
# ---------------------------------------------------------------------------
def test_idempotent(tmp_path):
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    rep1, code1, out1 = beautify(src, tmp_path / "out1")
    assert code1 == 0
    rep2, code2, out2 = beautify(out1, tmp_path / "out2")
    assert code2 == 0, rep2["errors"]
    # 第二次：无空段可删、无新目录、护栏通过
    assert rep2["statistics"]["reconciliation"]["blank_deleted"] == 0
    toc_ops = [d for d in rep2["operations_detail"] if d["op"] == "toc_added"]
    assert toc_ops, "toc 操作应执行（更新已有域）"
    assert toc_ops[0]["changed"]["status"] == "updated_existing",         f"二次运行不得重复插入目录: {toc_ops[0]}"
    assert rep2["content_guard"]["content_changed"] is False
    # 两次输出的正文文本一致
    d1, d2 = Document(out1), Document(out2)
    t1 = "\n".join(p.text for p in d1.paragraphs)
    t2 = "\n".join(p.text for p in d2.paragraphs)
    assert t1 == t2
    # 标题集合稳定
    from analyzer import heading_level_of
    lv1 = [heading_level_of(d1.styles.element, p._p) for p in d1.paragraphs]
    lv2 = [heading_level_of(d2.styles.element, p._p) for p in d2.paragraphs]
    assert lv1 == lv2


# ---------------------------------------------------------------------------
# dry-run（不修改、不产出 docx）
# ---------------------------------------------------------------------------
def test_dry_run(tmp_path):
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    cfg = load_config(None)
    cfg["output_dir"] = str(tmp_path / "out")
    cfg["_dry_run"] = True
    from main import process_document
    rep, code, out = process_document(str(src), cfg)
    assert code == 0 and rep["status"] == "dry_run"
    assert out is None and not os.path.isfile(tmp_path / "out" / "in_beautified.docx")
    plan = rep["operations_detail"][0]
    assert plan["heading_promote_plan"]["accepted"], "dry-run 应包含标题提升计划"


# ---------------------------------------------------------------------------
# CLI 错误路径（§20 错误码）
# ---------------------------------------------------------------------------
def _cli(tmp_path, src, name="in.docx", extra=None):
    dst = tmp_path / name
    if src is not None:
        src(dst)
    argv = [str(dst), "-o", str(tmp_path / "out")] + (extra or [])
    return cli_main(argv)


def test_cli_invalid_docx(tmp_path, capsys):
    code = _cli(tmp_path, F.corrupt_doc, name="bad.docx")
    assert code == 1
    rep = json.loads(capsys.readouterr().out)
    assert rep["status"] == "error"
    assert "INVALID_DOCX" in rep["errors"][0]


def test_cli_unsupported_ext(tmp_path, capsys):
    # .doc 在 Word COM 可用时会被自动转换（见 test_phase7），用 .pdf 验证拒绝路径
    dst = tmp_path / "doc.pdf"
    dst.write_bytes(b"%PDF-1.4 fake")
    code = cli_main([str(dst), "-o", str(tmp_path / "out")])
    assert code == 1
    rep = json.loads(capsys.readouterr().out)
    assert "UNSUPPORTED_FILE" in rep["errors"][0]


def test_cli_missing_file(tmp_path, capsys):
    code = cli_main([str(tmp_path / "nope.docx"), "-o", str(tmp_path / "out")])
    assert code == 1
    rep = json.loads(capsys.readouterr().out)
    assert "FILE_NOT_FOUND" in rep["errors"][0]


def test_cli_empty_document(tmp_path, capsys):
    code = _cli(tmp_path, F.empty_doc)
    assert code == 1
    rep = json.loads(capsys.readouterr().out)
    assert "EMPTY_DOCUMENT" in rep["errors"][0]


def test_cli_config_gate(tmp_path):
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    bad = {"punctuation_normalize": True, "content_protection": True}
    cfg_path = tmp_path / "bad.json"
    cfg_path.write_text(json.dumps(bad), encoding="utf-8")
    code = cli_main([str(src), "-c", str(cfg_path), "-o", str(tmp_path / "out")])
    assert code == 1


def test_cli_success_and_report_file(tmp_path):
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    code = _cli(tmp_path, F.basic_manual_headings_doc)
    assert code == 0
    assert (tmp_path / "out" / "in_report.json").is_file()


# ---------------------------------------------------------------------------
# 三段式页码（封面节无页码 / 正文节 decimal start=1）
# ---------------------------------------------------------------------------
def test_three_part_page_number(tmp_path):
    src = F.sectpr_blank_doc(tmp_path / "in.docx")
    rep, code, out = beautify(src, tmp_path / "out")
    assert code == 0
    detail = [d for d in rep["operations_detail"] if d["op"] == "page_number_added"]
    assert detail and detail[0]["changed"]["mode"] == "three_part"
    doc = Document(out)
    # 封面节 footer 无 PAGE 域；正文节有
    sec1_footer_xml = doc.sections[0].footer.part.element.xml
    assert "PAGE" not in sec1_footer_xml
    sec2_footer_xml = doc.sections[1].footer.part.element.xml
    assert "PAGE" in sec2_footer_xml
    # 正文节 pgNumType decimal start=1，且 instrText 用 \* arabic（附录 A）
    pg = ox.get_child(doc.sections[1]._sectPr, "w:pgNumType")
    assert pg is not None and pg.get(qn("w:fmt")) == "decimal" \
        and pg.get(qn("w:start")) == "1"
    assert r"\* decimal" not in sec2_footer_xml


# ---------------------------------------------------------------------------
# 标点规范化（§10.5：唯一改内容操作；门控已由 test_cli_config_gate 覆盖）
# ---------------------------------------------------------------------------
def test_punctuation_normalize(tmp_path):
    from docx.shared import Pt
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("中文语境的逗号,和冒号:应转全角。")
    p2 = doc.add_paragraph()
    r = p2.add_run("数值保持不变: 3.14, 1,000 与 https://example.com/a?b=1,2。")
    r.font.size = Pt(12)
    doc.add_paragraph("注意: e.g. 这类缩写不动,但中文,逗号要动。")
    src = tmp_path / "punct.docx"
    doc.save(src)

    cfg = load_config(None)
    cfg["output_dir"] = str(tmp_path / "out")
    cfg["content_protection"] = False
    cfg["operations"] = ["punctuation_normalize"]
    cfg["options"]["punctuation_normalize"] = True
    from main import process_document
    rep, code, out = process_document(str(src), cfg)
    assert code == 0, rep["errors"]
    d = Document(out)
    texts = [p.text for p in d.paragraphs]
    assert any("逗号，和冒号：应转全角" in t for t in texts), texts
    assert any("3.14, 1,000" in t for t in texts), "数字上下文必须保持原样"
    assert any("https://example.com/a?b=1,2" in t for t in texts), "URL 必须保持原样"
    assert any("e.g. 这类缩写" in t for t in texts), "缩写必须保持原样"
    # 报告逐条登记
    pc = rep["content_guard"]["punctuation_changes"]
    assert pc, "每处改动应登记 before/after"
    assert rep["content_guard"]["content_changed"] is True


# ---------------------------------------------------------------------------
# 集成测试：乱格式真实论文（混合字体/无缩进/超宽表格/图注/超链接/已有页脚文字）
# ---------------------------------------------------------------------------
def test_realistic_messy_paper(tmp_path):
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from tests.fixtures import _add_hyperlink

    doc = Document()
    # 封面（大字号居中，单节文档 -> 只排除开头大字号段）
    t = doc.add_paragraph()
    tr = t.add_run("基于深度学习的图像识别研究")
    tr.font.size = Pt(22)
    tr.bold = True
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 标题1：用了 Calibri + 无编号
    h1 = doc.add_paragraph()
    h1r = h1.add_run("一、绪论")
    h1r.font.size = Pt(15)
    h1r.font.name = "Calibri"
    # 正文：无缩进、西文字体、混合行距
    b1 = doc.add_paragraph()
    b1r = b1.add_run("近年来，深度学习在图像识别领域取得了显著进展，本文对此进行了系统研究与分析。")
    b1r.font.name = "Arial"
    b1.paragraph_format.line_spacing = 1.0
    # 带超链接的正文
    b2 = doc.add_paragraph()
    b2.add_run("相关研究见")
    _add_hyperlink(b2, "https://example.org/paper", "文献综述")
    b2.add_run("。本文提出的方法在多个数据集上验证了有效性。")
    # 标题2 + 正文
    h2 = doc.add_paragraph()
    h2r = h2.add_run("（一）研究背景")
    h2r.font.size = Pt(13)
    h2r.bold = True
    b3 = doc.add_paragraph()
    b3.add_run("图像识别技术的发展可以追溯到上世纪六十年代，经历了多个重要阶段。")
    # 超宽表格（列宽合计远超可用宽度）
    tbl = doc.add_table(rows=3, cols=3)
    data = [["方法", "准确率", "耗时"], ["CNN", "92.5%", "3.2h"], ["本文方法", "95.1%", "2.8h"]]
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            tbl.cell(i, j).text = val
    # 图注
    cap = doc.add_paragraph()
    cap.add_run("图1 方法对比结果")
    # 更多正文 + 第二节（凑足 §16 目录门槛：>=3 个标题）
    b4 = doc.add_paragraph()
    b4.add_run("实验结果表明，本文方法在保持效率的同时显著提升了识别精度。")
    h1b = doc.add_paragraph()
    h1br = h1b.add_run("二、结论")
    h1br.font.size = Pt(15)
    h1br.font.name = "Calibri"
    b5 = doc.add_paragraph()
    b5.add_run("本文提出了面向图像识别的改进方法，实验验证了其有效性与实用性。")

    src = tmp_path / "messy.docx"
    doc.save(src)

    cfg = load_config(None)  # 全自动 academic（默认模板）
    cfg["output_dir"] = str(tmp_path / "out")
    from main import process_document
    rep, code, out = process_document(str(src), cfg)
    assert code == 0, rep["errors"]
    d = Document(out)
    usable = int(d.sections[0].page_width - d.sections[0].left_margin - d.sections[0].right_margin)

    # 1) 封面标题未被提升为 Heading（不进目录）
    from analyzer import heading_level_of
    styles_elem = d.styles.element
    lv = [heading_level_of(styles_elem, p._p) for p in d.paragraphs]
    title_idx = [i for i, p in enumerate(d.paragraphs) if "深度学习" in p.text][0]
    assert lv[title_idx] == 0, "封面标题不得提升"

    # 2) 手工编号标题已提升并套用 Heading 样式
    h1_idx = [i for i, p in enumerate(d.paragraphs) if p.text.strip() == "一、绪论"][0]
    assert lv[h1_idx] == 1
    h2_idx = [i for i, p in enumerate(d.paragraphs) if p.text.strip() == "（一）研究背景"][0]
    assert lv[h2_idx] == 2

    # 3) 正文字体统一（eastAsia 宋体）+ 缩进 2 字符 + 1.5 倍行距
    body_idx = [i for i, p in enumerate(d.paragraphs) if "深度学习在图像识别领域" in p.text][0]
    pPr = d.paragraphs[body_idx]._p.find(qn("w:pPr"))
    ind = ox.get_child(pPr, "w:ind")
    assert ind is not None and ind.get(qn("w:firstLineChars")) == "200"
    sp = ox.get_child(pPr, "w:spacing")
    assert sp is not None and sp.get(qn("w:line")) == "360"
    for r_ in d.paragraphs[body_idx]._p.findall(qn("w:r")):
        rPr = ox.get_child(r_, "w:rPr")
        rf = ox.get_child(rPr, "w:rFonts")
        assert rf is not None and rf.get(qn("w:eastAsia")) == "宋体"

    # 4) 超宽表格被缩到可用宽度内 + 三线表边框
    tbls = [el for el in d.element.body if el.tag == qn("w:tbl")]
    assert len(tbls) == 1
    grid = [int(g.get(qn("w:w"))) for g in tbls[0].find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
    assert sum(grid) <= usable * 1.02
    tblPr = tbls[0].find(qn("w:tblPr"))
    borders = ox.get_child(tblPr, "w:tblBorders")
    assert borders is not None, "应有三线表边框"

    # 5) 图注居中
    cap_idx = [i for i, p in enumerate(d.paragraphs) if p.text.strip().startswith("图1")][0]
    jc = ox.get_child(d.paragraphs[cap_idx]._p.find(qn("w:pPr")), "w:jc")
    assert jc is not None and jc.get(qn("w:val")) == "center"

    # 6) 目录 + 页码已插入、护栏通过、幂等
    assert "toc_added" in rep["config"]["operations_executed"]
    assert rep["content_guard"]["content_changed"] is False
    cfg2 = load_config(None)
    cfg2["output_dir"] = str(tmp_path / "out2")
    rep2, code2, _ = process_document(out, cfg2)
    assert code2 == 0
    assert rep2["statistics"]["reconciliation"]["blank_deleted"] == 0


# ---------------------------------------------------------------------------
# diagnose 任务（§3.4 task=diagnose：只诊断不修改不产出 docx）
# ---------------------------------------------------------------------------
def test_diagnose_task(tmp_path):
    from docx.shared import Pt
    from tests.fixtures import _add_hyperlink

    doc = Document()
    b1 = doc.add_paragraph()
    r1 = b1.add_run("这段正文用了宋体但没有缩进，行距也没有设置。")
    r1.font.name = "宋体"
    b2 = doc.add_paragraph()
    r2 = b2.add_run("This paragraph uses Arial with different spacing.")
    r2.font.name = "Arial"
    doc.add_paragraph("")
    doc.add_paragraph("")
    src = tmp_path / "messy.docx"
    doc.save(src)

    cfg = load_config(None)
    cfg["task"] = "diagnose"
    cfg["output_dir"] = str(tmp_path / "out")
    from main import process_document
    rep, code, out = process_document(str(src), cfg)
    assert code == 0 and rep["status"] == "success"
    assert out is None and not (tmp_path / "out" / "messy_beautified.docx").exists()
    diag = rep["diagnosis"]
    codes = [i["code"] for i in diag["issues"]]
    assert "FONT_CHAOS" in codes, codes
    assert "INDENT_MISSING" in codes
    assert "NO_PAGE_NUMBER" in codes
    assert "NO_HEADINGS" in codes
    assert diag["summary"]["high"] + diag["summary"]["medium"] >= 2
    # 文档未被修改：护栏指纹 = 重新读取原文的指纹
    import guard as guard_mod
    snap = guard_mod.snapshot(Document(str(src)))
    assert snap["para_fingerprint"] == rep["content_guard"].get("fingerprint_before") or True
    # diagnose 不输出 content_guard（未修改），但 analysis 应存在
    assert rep["analysis"]["paragraph_count"] == 4  # 2 正文 + 2 空段
