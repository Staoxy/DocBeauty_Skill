# -*- coding: utf-8 -*-
"""V1.5 Phase 0 测试：CLI 子命令 + 版本号 + config 新字段。"""
import json
import os
import subprocess
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from config import ConfigError, load_config  # noqa: E402
from main import main as cli_main, process_document  # noqa: E402

from tests import fixtures as F  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def test_version_is_1_5_0(capsys):
    assert cli_main(["--version"]) == 0
    assert capsys.readouterr().out.strip() == "1.5.0"


def test_config_mode_validation():
    # mode=template 缺 template 路径 -> CONFIG_INVALID
    with pytest.raises(ConfigError):
        load_config({"mode": "template"})
    with pytest.raises(ConfigError):
        load_config({"mode": "reference"})
    cfg = load_config({"mode": "template", "template": "t.docx"})
    assert cfg["mode"] == "template" and cfg["template"] == "t.docx"
    # intent 字段可传
    cfg = load_config({"mode": "prompt", "intent": {"style": "formal"}})
    assert cfg["intent"] == {"style": "formal"}


def test_analyze_subcommand(tmp_path, capsys):
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    code = cli_main(["analyze", str(src), "-o", str(tmp_path / "out")])
    out = capsys.readouterr().out
    assert code == 0
    rep = json.loads(out)
    assert rep["status"] == "success"
    assert rep["diagnosis"]["issues"], "analyze 应输出诊断项"
    assert not (tmp_path / "out" / "in_beautified.docx").exists(), "analyze 不得产出 docx"
    assert (tmp_path / "out" / "in_report.json").is_file()


def test_verify_subcommand(tmp_path, capsys):
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    # 先产出一个美化文档
    cfg = load_config(None)
    cfg["output_dir"] = str(tmp_path / "out")
    process_document(str(src), cfg)
    out_docx = tmp_path / "out" / "in_beautified.docx"
    assert out_docx.is_file()

    code = cli_main(["verify", str(out_docx), "-o", str(tmp_path / "vout")])
    out = capsys.readouterr().out
    assert code == 0
    rep = json.loads(out)
    assert rep["status"] == "success"
    statuses = {c["id"]: c["status"] for c in rep["quality_checks"]}
    assert statuses["Q02"] == "skip" and statuses["Q03"] == "skip", "standalone 应跳过 diff 类"
    assert statuses["Q01"] == "pass"
    # Q08 标题层级：美化文档应有标题且无跳级
    assert statuses["Q08"] in ("pass", "info")


def test_verify_detects_broken_hierarchy(tmp_path, capsys):
    """verify 能发现层级跳级：Q08 为 warning 级（§18 设计），不影响 exit code。"""
    from docx import Document
    doc = Document()
    doc.add_heading("一、章", level=1)
    doc.add_heading("1.1.1 节", level=3)  # 跳级
    doc.add_paragraph("正文内容")
    src = tmp_path / "jumpy.docx"
    doc.save(src)
    code = cli_main(["verify", str(src), "-o", str(tmp_path / "out")])
    rep = json.loads(capsys.readouterr().out)
    assert code == 0
    q08 = [c for c in rep["quality_checks"] if c["id"] == "Q08"][0]
    assert q08["status"] == "warn" and "jumps" in q08["detail"]


def _com_available():
    try:
        import legacy_convert
        return legacy_convert.com_available()
    except Exception:
        return False


@pytest.mark.skipif(not _com_available(), reason="Word COM not available")
def test_render_subcommand(tmp_path, capsys):
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    cfg = load_config(None)
    cfg["output_dir"] = str(tmp_path / "out")
    process_document(str(src), cfg)
    out_docx = tmp_path / "out" / "in_beautified.docx"
    code = cli_main(["render", str(out_docx), "--pdf", "-o", str(tmp_path / "out")])
    out = capsys.readouterr().out
    assert code == 0
    rep = json.loads(out)
    rc = rep["render_check"]
    assert rc["available"] is True
    assert rc["page_count"] >= 1
    assert os.path.isfile(rc["pdf_exported"])


def test_render_without_com_is_error(tmp_path, capsys, monkeypatch):
    """COM 不可用时显式 render 请求 = error（不静默成功）。"""
    import render_check
    monkeypatch.setattr(render_check, "render_check",
                        lambda *a, **k: {"available": False, "reason": "mocked"})
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    code = cli_main(["render", str(src), "-o", str(tmp_path / "out")])
    rep = json.loads(capsys.readouterr().out)
    assert code == 1 and rep["status"] == "error"


def test_subprocess_apply_still_works(tmp_path):
    """子进程级回归：兼容模式与 apply 子命令等价。"""
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    for argv in ([str(src), "-o", str(tmp_path / "o1")],
                 ["apply", str(src), "-o", str(tmp_path / "o2")]):
        r = subprocess.run([sys.executable, "main.py"] + argv,
                           capture_output=True, encoding="utf-8", cwd=ROOT)
        assert r.returncode == 0, r.stderr[-300:]
        rep = json.loads(r.stdout)
        assert rep["status"] == "success"
