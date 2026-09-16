# -*- coding: utf-8 -*-
"""V1.5 Phase 3 测试：Intent 校验/映射 + prompt 模式。"""
import os
import sys

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import intent  # noqa: E402
from config import load_config  # noqa: E402
from main import process_document  # noqa: E402
from templates import get_template  # noqa: E402
from utils import ooxml as ox  # noqa: E402

from tests import fixtures as F  # noqa: E402


def test_validate_intent_clean():
    normalized, warnings = intent.validate_intent({
        "style": "formal",
        "document_type": "report",
        "preferences": {"line_density": "loose", "heading_emphasis": "strong"},
    })
    assert normalized["style"] == "formal"
    assert normalized["preferences"]["line_density"] == "loose"
    assert warnings == []


def test_validate_intent_degrades_gracefully():
    normalized, warnings = intent.validate_intent({
        "style": "super_fancy",              # 枚举外
        "unknown_key": 1,                     # 未知键
        "preferences": {"line_density": "huge", "mystery": True},
    })
    assert normalized["style"] is None
    assert "line_density" not in normalized["preferences"]
    assert len(warnings) >= 3, warnings
    # 非对象输入
    normalized, warnings = intent.validate_intent("排版好看点")
    assert normalized["style"] is None and warnings


def test_intent_to_params_mapping():
    builtin = get_template("formal")  # h1 before 18 / after 12, line 1.5
    normalized, _ = intent.validate_intent({
        "preferences": {"heading_emphasis": "strong", "line_density": "loose",
                        "table_density": "comfortable", "caption_align": "left"},
    })
    layer, echo = intent.intent_to_params(normalized, builtin)
    assert layer["headings"]["h1"]["before_pt"] == 27.0  # 18 * 1.5
    assert layer["paragraph"]["line_spacing"] == 1.75
    assert layer["tables"]["cell_margin_dxa"] == 160
    assert layer["captions"]["align"] == "left"
    assert echo["heading_emphasis"]["effect"]
    # default 偏好不产生参数
    normalized2, _ = intent.validate_intent({"preferences": {"line_density": "default"}})
    layer2, _ = intent.intent_to_params(normalized2, builtin)
    assert "paragraph" not in layer2


def test_prompt_mode_e2e(tmp_path):
    target = F.messy_paper_doc(tmp_path / "target.docx")
    cfg = load_config(None)
    cfg["mode"] = "prompt"
    cfg["output_dir"] = str(tmp_path / "out")
    cfg["intent"] = {
        "style": "formal",
        "content_protection": True,
        "preferences": {"line_density": "loose", "table_density": "comfortable",
                        "caption_align": "left", "heading_emphasis": "strong"},
    }
    rep, code, out = process_document(str(target), cfg)
    assert code == 0, rep["errors"]
    assert rep["content_guard"]["content_changed"] is False
    assert rep["config"]["style_resolved"] == "formal"
    assert rep["intent_resolved"]["line_density"]["value"] == "loose"

    d = Document(out)
    # 行距 1.75 → line=420
    body = [p for p in d.paragraphs if "深度学习在图像识别领域" in p.text][0]
    sp = ox.get_child(body._p.find(qn("w:pPr")), "w:spacing")
    assert sp.get(qn("w:line")) == "420"
    # 单元格边距 160
    tbl = [el for el in d.element.body if el.tag == qn("w:tbl")][0]
    mar = tbl.find(qn("w:tblPr") + "/" + qn("w:tblCellMar")) \
        if tbl.find(qn("w:tblPr") + "/" + qn("w:tblCellMar")) is not None \
        else tbl.find(".//" + qn("w:tblCellMar"))
    left = ox.get_child(mar, "w:left")
    assert left.get(qn("w:w")) == "160"
    # 题注左对齐
    cap = [p for p in d.paragraphs if p.text.strip().startswith("图1")][0]
    jc = ox.get_child(cap._p.find(qn("w:pPr")), "w:jc")
    assert jc.get(qn("w:val")) == "left"
    # 标题间距 strong：formal h1 before 18 → 27
    h1_style = ox.find_style_by_name(d.styles.element, "heading 1")
    sp2 = ox.get_child(ox.get_child(h1_style, "w:pPr"), "w:spacing")
    assert sp2.get(qn("w:before")) == "540"  # 27pt * 20


def test_prompt_without_intent_falls_back(tmp_path):
    target = F.messy_paper_doc(tmp_path / "target.docx")
    cfg = load_config(None)
    cfg["mode"] = "prompt"
    cfg["output_dir"] = str(tmp_path / "out")
    rep, code, out = process_document(str(target), cfg)
    assert code == 0
    assert any("intent" in w for w in rep["warnings"])


def test_intent_cannot_flip_protection_gate(tmp_path):
    target = F.messy_paper_doc(tmp_path / "target.docx")
    cfg = load_config(None)
    cfg["mode"] = "prompt"
    cfg["output_dir"] = str(tmp_path / "out")
    cfg["content_protection"] = True
    cfg["intent"] = {"content_protection": False}
    rep, code, out = process_document(str(target), cfg)
    assert code == 0
    assert any("content_protection" in w for w in rep["warnings"]), \
        "intent 翻转保护门控必须产生警告"
    # 门控保持开启：护栏仍启用
    assert rep["content_guard"]["enabled"] is True
