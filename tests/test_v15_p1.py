# -*- coding: utf-8 -*-
"""V1.5 Phase 1 测试：Style Analyzer + Template Analyzer 双路径。"""
import json
import os
import sys

import jsonschema
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import style_analyzer  # noqa: E402
import template_analyzer  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from utils import ooxml as ox  # noqa: E402

from tests import fixtures as F  # noqa: E402

_SPEC_SCHEMA = json.load(open(
    os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                 "schemas", "template_spec.json"), encoding="utf-8"))


def test_style_analyzer_regular(tmp_path):
    src = F.regular_template_doc(tmp_path / "tpl.docx")
    doc = Document(str(src))
    sa = style_analyzer.analyze_styles(doc)
    # docDefaults
    assert sa["doc_defaults"]["eastAsia"] in ("等线", "Calibri", None)  # 默认模板主题
    # Normal 样式
    normal = sa["styles"]["Normal"]
    assert normal["eastAsia"] == "仿宋"
    assert normal["ascii"] == "Georgia"
    assert normal["size_pt"] == 12.0
    # Heading 样式带 outlineLvl
    h1 = [v for v in sa["styles"].values() if v.get("outline_lvl") == 0]
    assert h1 and h1[0]["eastAsia"] == "黑体" and h1[0]["size_pt"] == 16.0
    # 使用直方图
    assert sa["style_usage"].get("Heading 1", 0) == 1
    # 声明 vs 实际统计存在
    assert sa["declared_vs_actual"]["total_body_paras"] >= 2


def test_template_analyzer_regular_path(tmp_path):
    src = F.regular_template_doc(tmp_path / "tpl.docx")
    doc = Document(str(src))
    spec = template_analyzer.analyze_template(doc)
    jsonschema.validate(spec, _SPEC_SCHEMA)
    assert spec["template_kind"] == "regular"
    # roles 来自样式定义
    assert spec["roles"]["body"]["eastAsia"] == "仿宋"
    assert spec["roles"]["body"]["ascii"] == "Georgia"
    assert spec["roles"]["h1"]["eastAsia"] == "黑体"
    assert spec["roles"]["h1"]["size_pt"] == 16.0
    assert spec["roles"]["h2"]["size_pt"] == 14.0
    # 页面与页眉页脚
    assert spec["page"]["size"] == "A4"
    assert spec["page"]["margins_cm"]["left"] == 3.0
    assert spec["header_footer"]["header"]["text"] == "XX大学课程论文"
    assert spec["header_footer"]["footer_page_number"]["present"] is True
    assert spec["header_footer"]["footer_page_number"]["align"] == "center"
    # roles_meta 可回溯来源
    assert spec["roles_meta"]["h1"] == "heading 1"


def test_template_analyzer_sample_path(tmp_path):
    src = F.sample_template_doc(tmp_path / "sample_tpl.docx")
    doc = Document(str(src))
    spec = template_analyzer.analyze_template(doc)
    jsonschema.validate(spec, _SPEC_SCHEMA)
    assert spec["template_kind"] == "sample"
    # 样例路径：从直接格式提取
    assert spec["roles"]["h1"]["eastAsia"] == "黑体"
    assert spec["roles"]["h1"]["size_pt"] == 16.0
    assert spec["roles"]["h2"]["size_pt"] == 14.0
    assert spec["roles"]["body"]["eastAsia"] == "宋体"
    assert spec["roles"]["body"]["size_pt"] == 12.0
    assert spec["roles"]["body"]["line"] == 1.5
    assert spec["roles"]["body"]["first_line_chars"] == 2.0
    assert spec["roles"]["caption"]["size_pt"] == 10.5
    assert spec["roles"]["title"]["size_pt"] == 22.0
    assert spec["roles_meta"]["h1"] == "1 paragraphs"
    assert spec["page"]["margins_cm"]["left"] == 3.0
    assert spec["header_footer"]["footer_page_number"]["present"] is True


def test_template_kind_decision(tmp_path):
    """判定逻辑：无样式标题且无直接格式标题层级 → 回退。"""
    doc = Document()
    doc.add_paragraph("只有一个普通段落。")
    doc.save(str(tmp_path / "empty_tpl.docx"))
    spec = template_analyzer.analyze_template(Document(str(tmp_path / "empty_tpl.docx")))
    assert spec["template_kind"] == "sample"
    assert "h1" not in spec["roles"] or not spec["roles"].get("h1")
    # spec 应可校验
    jsonschema.validate(spec, _SPEC_SCHEMA)


def test_template_unsupported_columns(tmp_path):
    """分栏模板 → unsupported 记录（边界诚实）。"""
    from docx.enum.section import WD_SECTION
    doc = Document()
    doc.add_heading("一、标题", level=1)
    doc.add_paragraph("正文")
    sectPr = doc.sections[0]._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        from docx.oxml import OxmlElement
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), "2")
    doc.save(str(tmp_path / "cols_tpl.docx"))
    spec = template_analyzer.analyze_template(Document(str(tmp_path / "cols_tpl.docx")))
    assert "columns" in spec["unsupported"]
