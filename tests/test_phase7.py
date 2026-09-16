# -*- coding: utf-8 -*-
"""Phase 7 测试：旧格式转换 + 渲染级验证。

需要真实 Word COM（Windows + Office）。COM 不可用时自动 skip——
这是环境增强项，不是核心功能门槛。
"""
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from config import load_config  # noqa: E402
from main import main as cli_main  # noqa: E402

from tests import fixtures as F  # noqa: E402


def _com_available() -> bool:
    try:
        import legacy_convert
        return legacy_convert.com_available()
    except Exception:
        return False


COM = _com_available()
requires_word = pytest.mark.skipif(not COM, reason="Word COM not available")


def _make_doc_file(path):
    """用 Word COM 把 docx fixture 另存为 .doc（wdFormatDocument97=0）。"""
    import pythoncom
    import win32com.client
    from docx import Document
    Document().save(str(path) + ".tmp.docx")
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("Word.Application")
    app.Visible = False
    app.DisplayAlerts = 0
    doc = None
    try:
        doc = app.Documents.Open(os.path.abspath(str(path) + ".tmp.docx"),
                                 AddToRecentFiles=False, Visible=False)
        doc.Content.Text = "一、标题\r\n这是正文内容，用于验证旧格式转换。\r\n二、结论\r\n正文结束。"
        doc.SaveAs2(os.path.abspath(str(path)), FileFormat=0)
    finally:
        try:
            if doc is not None:
                doc.Close(SaveChanges=0)
        except Exception:
            pass
        app.Quit()
    return str(path)


@requires_word
def test_legacy_convert_doc_to_docx(tmp_path):
    import legacy_convert
    src = _make_doc_file(tmp_path / "legacy.doc")
    dst = tmp_path / "legacy_converted.docx"
    info = legacy_convert.convert_to_docx(src, str(dst))
    assert os.path.isfile(dst)
    from docx import Document
    d = Document(str(dst))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "一、标题" in text and "正文内容" in text


@requires_word
def test_cli_accepts_doc_input(tmp_path, capsys):
    import json
    src = _make_doc_file(tmp_path / "报告.doc")
    code = cli_main([src, "-o", str(tmp_path / "out")])
    out = capsys.readouterr().out
    assert code == 0, out
    rep = json.loads(out)
    assert rep["status"] == "success"
    assert "converted_from" in rep["input"]
    assert (tmp_path / "out" / "报告_converted.docx").is_file()


@requires_word
def test_render_check_full_pipeline(tmp_path):
    """端到端：beautify + --render-check --pdf：真 Word 验证目录实体化与页码。"""
    import json as _json
    import subprocess
    src = F.basic_manual_headings_doc(tmp_path / "in.docx")
    outdir = tmp_path / "out"
    r = subprocess.run(
        [sys.executable, "main.py", str(src), "-o", str(outdir),
         "--render-check", "--pdf"],
        capture_output=True, encoding="utf-8", cwd=os.path.dirname(
            os.path.dirname(os.path.abspath(__file__))))
    assert r.returncode == 0, r.stderr[-500:]
    rep = _json.loads(r.stdout)
    rc = rep.get("render_check", {})
    assert rc.get("available") is True, rc
    assert rc.get("page_count", 0) >= 1
    toc = rc.get("toc", {})
    assert toc.get("tables_of_contents") == 1
    assert toc.get("has_real_entries") is True, toc
    assert toc.get("entries", 0) >= 3, "目录应有 >=3 条真实条目"
    assert rc.get("page_number_rendered") is True
    assert "绪论" in (toc.get("preview") or "") or "方法" in (toc.get("preview") or "")
    assert os.path.isfile(rc.get("pdf_exported", "")), "应导出 PDF"


@requires_word
def test_render_check_unavailable_graceful(tmp_path):
    """render_check 对不存在文件优雅返回 available=False，不抛异常。"""
    import render_check
    rc = render_check.render_check(str(tmp_path / "nope.docx"))
    assert rc.get("available") is False
