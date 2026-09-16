# -*- coding: utf-8 -*-
"""V1.5 Phase 5 测试：visual_check 客观指标。"""
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import visual_check  # noqa: E402
from config import load_config  # noqa: E402
from main import main as cli_main, process_document  # noqa: E402

from tests import fixtures as F  # noqa: E402


def _make_pdf(path, with_blank=True, with_overflow=False):
    """直接用 pymupdf 构造测试 PDF（不依赖 Word）。

    注意 insert_text 不换行：每行单独调用，行宽控制在边距内。
    """
    import pymupdf
    doc = pymupdf.open()
    # 第 1 页：8 行正常正文（行宽 ~240pt，页宽 595pt、边距 72pt 内）
    page = doc.new_page(width=595, height=842)  # A4 pt
    for i in range(8):
        page.insert_text((72, 100 + i * 40),
                         f"Line {i + 1}: body text for the first page content.",
                         fontsize=12)
    if with_blank:
        doc.new_page(width=595, height=842)  # 空白页
    page3 = doc.new_page(width=595, height=842)
    for i in range(8):
        page3.insert_text((72, 100 + i * 40),
                          f"Line {i + 1}: third page body content here.",
                          fontsize=12)
    if with_overflow:
        # 越界矩形：A4 宽 595pt、右边距 72pt → 右边界 700 越出页面
        page3.draw_rect(pymupdf.Rect(50, 600, 700, 700), color=(0, 0, 0), width=1)
    doc.save(str(path))
    return str(path)


def test_visual_blank_and_overflow(tmp_path):
    pdf = _make_pdf(tmp_path / "t.pdf", with_blank=True, with_overflow=True)
    result = visual_check.check_pdf(pdf, margins_cm={"left": 2.54, "right": 2.54},
                                    png_dir=str(tmp_path / "png"))
    assert result["available"] is True
    assert result["page_count"] == 3
    pages = result["pages"]
    assert "blank" in pages[1]["flags"], pages[1]
    assert (tmp_path / "png" / "page_002.png").is_file()
    issues = result["issues"]
    codes = [i["code"] for i in issues]
    assert "blank" in codes
    assert "out_of_margin" in codes, issues
    assert result["summary"]["result"] == "ERROR"
    assert result["summary"]["errors"] >= 2


def test_visual_clean_pass(tmp_path):
    pdf = _make_pdf(tmp_path / "clean.pdf", with_blank=False, with_overflow=False)
    result = visual_check.check_pdf(pdf, margins_cm={"left": 2.0, "right": 2.0})
    assert result["available"] is True
    codes = [i["code"] for i in result["issues"]]
    assert "blank" not in codes
    assert "out_of_margin" not in codes
    assert result["summary"]["result"] == "PASS"


def _com_available():
    try:
        import legacy_convert
        return legacy_convert.com_available()
    except Exception:
        return False


@pytest.mark.skipif(not _com_available(), reason="Word COM not available")
def test_render_visual_e2e_blank_page(tmp_path):
    """端到端：含空白页的文档 → render --pdf → visual 标记 blank。"""
    from docx import Document
    doc = Document()
    doc.add_heading("一、标题", level=1)
    for i in range(3):
        doc.add_paragraph(f"正文内容第{i + 1}段，用于填充第一页的正常内容展示。")
    doc.add_page_break()  # 产生一个几乎空白的页（cleanup 会保护含分页符的段落）
    doc.save(str(tmp_path / "with_blank.docx"))

    code = cli_main(["render", str(tmp_path / "with_blank.docx"), "--pdf",
                     "-o", str(tmp_path / "out")])
    assert code == 0
    # 直接检查产物（stdout 属于 render 报告）
    import json
    import subprocess
    r = subprocess.run([sys.executable, "main.py", "render",
                        str(tmp_path / "with_blank.docx"), "--pdf",
                        "-o", str(tmp_path / "out2")],
                       capture_output=True, encoding="utf-8",
                       cwd=os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    rep = json.loads(r.stdout)
    visual = rep["render_check"].get("visual", {})
    assert visual.get("available") is True
    codes = [i["code"] for i in visual.get("issues", [])]
    assert "blank" in codes, visual.get("issues")
    assert os.path.isfile(rep["render_check"]["pdf_exported"])
