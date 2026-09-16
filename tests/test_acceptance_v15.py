# -*- coding: utf-8 -*-
"""V1.5 十三步验收测试（DESIGN_V15 §19 完成定义）。

输入：我的论文.docx（乱格式）+ 学校论文模板.docx（样例型）
要求："按照学校模板排版，不要修改正文内容。"

步骤映射：
  ①②读取/识别结构   → 流水线内建
  ③分析模板          → template_spec
  ④建立 Role Map     → role_map
  ⑤Formatting Spec   → spec_builder（经 plan 可见）
  ⑥修改计划          → plan 子命令（dry_run + formatting_spec.json）
  ⑦执行排版          → apply
  ⑧内容未变          → content_guard
  ⑨DOCX 完整         → checker Q01
  ⑩渲染 PDF          → render（需 Word COM，skipif）
  ⑪视觉布局          → visual 指标（需 Word COM，skipif）
  ⑫最终 Word         → 产物存在
  ⑬report.json       → 产物存在
"""
import json
import os
import subprocess
import sys

import pytest
from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from config import load_config  # noqa: E402
from main import process_document  # noqa: E402
from utils import ooxml as ox  # noqa: E402

from tests import fixtures as F  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def test_acceptance_core(tmp_path):
    """①–⑨ + ⑫⑬：核心验收（无 Word 依赖）。"""
    target = F.messy_paper_doc(tmp_path / "我的论文.docx")
    template = F.sample_template_doc(tmp_path / "学校论文模板.docx")
    outdir = tmp_path / "output"

    # ⑥ 修改计划（plan，不执行）
    r = subprocess.run(
        [sys.executable, "main.py", "plan", str(target), "--template", str(template),
         "-o", str(outdir)],
        capture_output=True, encoding="utf-8", cwd=ROOT)
    assert r.returncode == 0, r.stderr[-400:]
    plan = json.loads(r.stdout)
    assert plan["status"] == "dry_run"
    # ③④⑤
    assert plan["template_spec"]["template_kind"] == "sample"
    assert plan["role_map"]["body"], "sample 模板的 body 角色应被识别"
    spec = json.loads((outdir / "formatting_spec.json").read_text(encoding="utf-8"))
    assert spec["mode"] == "template"
    assert spec["params"]["font"]["chinese"] == "宋体"
    assert spec["sources"]["font.chinese"] == "P2:template"
    assert not (outdir / "我的论文_beautified.docx").exists(), "plan 阶段不得产出"

    # ⑦ 执行排版
    cfg = load_config(None)
    cfg["mode"] = "template"
    cfg["template"] = str(template)
    cfg["output_dir"] = str(outdir)
    rep, code, out_docx = process_document(str(target), cfg)
    assert code == 0, rep["errors"]

    # ⑧ 内容未变
    assert rep["content_guard"]["content_changed"] is False
    assert rep["content_guard"]["enabled"] is True
    # ⑨ DOCX 完整
    q01 = [c for c in rep["quality_checks"] if c["id"] == "Q01"][0]
    assert q01["status"] == "pass"
    # ⑫ 最终 Word 存在且格式符合模板
    assert out_docx and os.path.isfile(out_docx)
    d = Document(out_docx)
    body = [p for p in d.paragraphs if "深度学习在图像识别领域" in p.text][0]
    rf = ox.get_child(ox.get_child(body._p.findall(qn("w:r"))[0], "w:rPr"), "w:rFonts")
    assert rf.get(qn("w:eastAsia")) == "宋体"
    pPr = body._p.find(qn("w:pPr"))
    ind = ox.get_child(pPr, "w:ind")
    assert ind.get(qn("w:firstLineChars")) == "200"
    assert round(d.sections[0].left_margin.cm, 2) == 3.0
    header_text = "".join(p.text for p in d.sections[0].header.paragraphs)
    assert "XX大学课程论文" in header_text
    assert "toc_added" in rep["config"]["operations_executed"]
    # ⑬ report.json
    report_path = outdir / "我的论文_report.json"
    assert report_path.is_file()
    saved = json.loads(report_path.read_text(encoding="utf-8"))
    assert saved["status"] == "success"
    assert saved["content_guard"]["content_changed"] is False


def _com_available():
    try:
        import legacy_convert
        return legacy_convert.com_available()
    except Exception:
        return False


@pytest.mark.skipif(not _com_available(), reason="Word COM not available")
def test_acceptance_visual(tmp_path):
    """⑩ 渲染 PDF + ⑪ 视觉布局（需 Word COM）。"""
    target = F.messy_paper_doc(tmp_path / "我的论文.docx")
    template = F.sample_template_doc(tmp_path / "学校论文模板.docx")
    outdir = tmp_path / "output"
    cfg = load_config(None)
    cfg["mode"] = "template"
    cfg["template"] = str(template)
    cfg["output_dir"] = str(outdir)
    cfg["_render_check"] = True
    cfg["_pdf"] = True
    rep, code, out_docx = process_document(str(target), cfg)
    assert code == 0, rep["errors"]
    rc = rep["render_check"]
    assert rc["available"] is True
    assert os.path.isfile(rc["pdf_exported"]), "⑩ 渲染 PDF"
    assert rc["toc"]["has_real_entries"] is True, "目录实体化"
    assert rc["page_number_rendered"] is True
    # ⑪ 视觉布局
    visual = rc.get("visual") or {}
    assert visual.get("available") is True
    assert visual["summary"]["result"] in ("PASS", "WARNING"), visual["issues"]
    assert os.path.isdir(outdir / "render_pages")
    # 完成判定（§30）
    assert rep["status"] == "success"
