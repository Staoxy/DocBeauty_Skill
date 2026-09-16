# -*- coding: utf-8 -*-
"""DocBeauty V2 主入口（DESIGN_V2.md §5/§20/§21）。

CLI:
  python main.py <input.docx> [-c config.json] [-o OUTDIR] [--report PATH]
                 [--dry-run] [--batch DIR] [--print-schema]

退出码：0 成功 | 2 部分成功 | 3 护栏失败（不产出文件）| 1 错误。
stdout 只输出 report JSON；人类日志走 stderr（§21.1）。
"""
import argparse
import glob
import hashlib
import json
import os
import shutil
import sys
import tempfile
import time
import traceback
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import constants as C
from config import ConfigError, build_effective, load_config, load_schema, output_paths


def LOG(msg):
    print(f"[docbeauty] {msg}", file=sys.stderr)


class OpSkip(Exception):
    def __init__(self, reason):
        self.reason = reason


# ---------------------------------------------------------------------------
# 输入校验（§3.1）
# ---------------------------------------------------------------------------
def validate_input(path: str):
    if not os.path.isfile(path):
        return C.FILE_NOT_FOUND, f"input not found: {path}"
    if os.path.splitext(path)[1].lower() != ".docx":
        return C.UNSUPPORTED_FILE, "only .docx is supported (.doc/.wps/pdf 等请先转换)"
    if os.path.getsize(path) > C.MAX_FILE_BYTES:
        return C.FILE_TOO_LARGE, f"file exceeds {C.MAX_FILE_BYTES // (1024 * 1024)}MB limit"
    try:
        with zipfile.ZipFile(path) as zf:
            bad = zf.testzip()
            if bad is not None:
                return C.INVALID_DOCX, f"corrupt zip member: {bad}"
            names = zf.namelist()
            if "word/document.xml" not in names:
                return C.INVALID_DOCX, "missing word/document.xml"
            try:
                zf.read(names[0])
            except RuntimeError:
                return C.PASSWORD_PROTECTED, "encrypted docx is not supported"
    except zipfile.BadZipFile:
        return C.INVALID_DOCX, "not a valid zip/docx file"
    return None, None


def sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------------------
# 单文档流水线（§5.2 固定顺序）
# ---------------------------------------------------------------------------
def process_document(input_path: str, cfg: dict) -> tuple:
    """返回 (report, exit_code, output_docx_path)。"""
    import analyzer
    import checker
    import cleanup
    import diagnose as diagnose_mod
    import formatter
    import guard
    import headings
    import images
    import layout
    import lists
    import punctuation
    import structure
    import tables
    import toc as toc_mod
    from docx import Document
    from report import build_report, write_report

    t0 = time.time()
    sections = {}
    docx_path, report_path = output_paths(cfg, input_path)
    sections["input"] = {
        "file": os.path.basename(input_path),
        "path": os.path.abspath(input_path),
        "size_bytes": os.path.getsize(input_path),
        "sha256": sha256_file(input_path),
        **(cfg.get("_converted_note") or {}),
    }
    sections["output"] = {"file": os.path.basename(docx_path),
                          "path": os.path.abspath(docx_path), "written": False}

    def finalize(status, code):
        sections.setdefault("errors", [])
        sections.setdefault("warnings", [])
        sections["timing"] = {"total_ms": int((time.time() - t0) * 1000)}
        return build_report(status, code, sections)

    # ---- 1. 加载 + 护栏基线快照（任何修改之前） ----
    workdir = tempfile.mkdtemp(prefix="docbeauty_", dir=C.temp_dir_root())
    try:
        work_file = os.path.join(workdir, "work.docx")
        shutil.copy2(input_path, work_file)
        doc = Document(work_file)

        before_guard = guard.snapshot(doc)
        before_analysis = analyzer.analyze(doc)
        if before_analysis["paragraph_count"] == 0 and before_analysis["table_count"] == 0:
            raise ConfigError(C.EMPTY_DOCUMENT, "document has no paragraphs and no tables")

        # ---- 2/3. 结构检测 + 列表登记 ----
        zones = structure.detect(doc, before_analysis)
        list_count = lists.detect_lists(doc, zones)
        before_analysis["list_paragraph_count"] = list_count
        if zones.zones.get("toc_existing_manual"):
            before_analysis["existing_manual_toc"] = True

        # ---- 4. 配置合并（V1.5：template/reference 模式走 Rule Resolver） ----
        op_detail, executed, skipped, failed = [], [], [], []
        warnings = []
        accounting = {}
        if cfg["document_type"] == "auto":
            dtype = analyzer.detect_document_type(doc, before_analysis)
        else:
            dtype = {"value": cfg["document_type"], "confidence": 1.0, "evidence": []}

        import spec_builder
        import template_analyzer
        template_spec = reference_profile = None
        intent_normalized = None
        if cfg["mode"] == "prompt":
            import intent as intent_mod
            if cfg.get("intent") is None:
                warnings.append("mode=prompt 未提供 intent，回退 auto 模式")
                cfg = dict(cfg)
                cfg["mode"] = "auto"
            else:
                intent_normalized, intent_warnings = intent_mod.validate_intent(cfg["intent"])
                warnings.extend(intent_warnings)
                if intent_normalized["style"]:
                    cfg = dict(cfg)
                    cfg["style"] = intent_normalized["style"]  # P1: 意图选择内置基线
                if intent_normalized["document_type"] and cfg["document_type"] == "auto":
                    dtype = {"value": intent_normalized["document_type"],
                             "confidence": 1.0, "evidence": ["intent.document_type"]}
        if cfg["mode"] in ("template", "reference"):
            src_path = cfg["template"] if cfg["mode"] == "template" else cfg["reference"]
            verr, vmsg = validate_input(src_path)
            if verr:
                raise ConfigError(C.CONFIG_INVALID,
                                  f"{cfg['mode']} file invalid: {vmsg}")
            from docx import Document as _TplDoc
            analyzed = template_analyzer.analyze_template(_TplDoc(src_path))
            if cfg["mode"] == "template":
                template_spec = analyzed
            else:
                reference_profile = analyzed

        from templates import get_template
        from config import resolve_style
        style_resolved = resolve_style(cfg, dtype["value"])
        if cfg["mode"] in ("template", "reference", "prompt"):
            intent_params = None
            intent_echo = {}
            if cfg["mode"] == "prompt" and intent_normalized is not None:
                import intent as intent_mod
                intent_params, intent_echo = intent_mod.intent_to_params(
                    intent_normalized, get_template(style_resolved))
                sections["intent_resolved"] = intent_echo
                if intent_normalized["content_protection"] is not None and                         intent_normalized["content_protection"] != cfg["content_protection"]:
                    warnings.append(
                        "intent.content_protection 不改变内容保护门控；请通过顶层字段设置")
            spec = spec_builder.build_formatting_spec(
                cfg, get_template(style_resolved), template_spec=template_spec,
                reference_profile=reference_profile,
                target_analysis=before_analysis, intent_params=intent_params)
            effective = spec["params"]
            if template_spec is not None:
                label = f"template:{template_spec['template_kind']}"
            elif cfg["mode"] == "reference":
                label = f"reference:{style_resolved}"
            else:
                label = style_resolved
            effective["_style_resolved"] = label
            sections["template_spec"] = template_spec
            sections["role_map"] = spec["role_map"]
            sections["formatting_spec"] = {
                "sources": spec["sources"], "policy": spec["policy"],
                "unsupported": spec["unsupported"]}
            warnings.extend(spec["warnings"])
        else:
            effective = build_effective(cfg, dtype["value"])

        ops = cfg["operations"] if isinstance(cfg["operations"], list) else C.AUTO_OPERATIONS
        sections["config"] = {
            "style_resolved": effective["_style_resolved"],
            "document_type": dtype,
            "content_protection": cfg["content_protection"],
            "operations_requested": cfg["operations"],
            "operations_planned": [C.OPERATION_CATALOG.get(o, o) for o in ops],
            "overrides_applied": cfg.get("overrides") or {},
            "options": cfg["options"],
        }

        # ---- task=diagnose：只诊断不修改（§3.4），不产出 docx ----
        if cfg.get("task") == "diagnose":
            diagnosis = diagnose_mod.diagnose(doc, before_analysis, zones)
            sections["diagnosis"] = diagnosis
            sections["analysis"] = before_analysis
            sections["zones"] = zones.to_report()
            sections["config"].pop("operations_planned", None)
            sections["config"]["operations_executed"] = ["diagnosis_done"]
            rep = finalize("success", C.EXIT_SUCCESS)
            write_report(rep, report_path)
            return rep, C.EXIT_SUCCESS, None

        # ---- dry-run：不修改任何东西（§21.2） ----
        if cfg.get("_dry_run"):
            sections["zones"] = zones.to_report()
            sections["analysis"] = before_analysis
            sections["operations_detail"] = [{
                "op": "dry_run_plan",
                "planned_operations": [C.OPERATION_CATALOG.get(o, o) for o in ops],
                "heading_promote_plan": {
                    "accepted": [{"paragraph": i, "level": lv,
                                  "text": analyzer.para_text(doc.paragraphs[i]._p).strip()[:40]}
                                 for i, lv in sorted(zones.headings.items())],
                    "candidates_not_modified": zones.candidates,
                },
                "toc_will_insert": ("toc" in ops
                                    and not before_analysis["existing_toc_field"]
                                    and len(zones.headings) >= 3),
            }]
            rep = finalize("dry_run", C.EXIT_SUCCESS)
            write_report(rep, report_path)
            if sections.get("formatting_spec"):
                spec_out = {
                    "mode": sections.get("formatting_spec", {}).get("mode", cfg["mode"]),
                    "params": {k: v for k, v in effective.items()
                               if not k.startswith("_")},
                    "role_map": sections.get("role_map", {}),
                    "sources": sections["formatting_spec"].get("sources", {}),
                    "policy": sections["formatting_spec"].get("policy", {}),
                    "unsupported": sections["formatting_spec"].get("unsupported", []),
                }
                with open(os.path.join(cfg["output_dir"], "formatting_spec.json"),
                          "w", encoding="utf-8") as sf:
                    json.dump(spec_out, sf, ensure_ascii=False, indent=2)
            return rep, C.EXIT_SUCCESS, None

        # ---- 5-15. 执行操作（容器已在步骤 4 初始化） ----
        if cfg["mode"] == "template" and "header_footer" not in ops:
            ops = list(ops) + ["header_footer"]  # 仅模板模式套用页眉文字（§3 MODE 4 参考模式不复制）

        def run_op(op, fn):
            if op not in ops:
                return None
            mapped = C.OPERATION_CATALOG.get(op, op)
            try:
                detail = fn() or {}
                warns = detail.pop("warnings", None) if isinstance(detail, dict) else None
                if warns:
                    warnings.extend(warns)
                executed.append(mapped)
                op_detail.append({"op": mapped, "status": "done", "changed": detail})
                return detail
            except OpSkip as s:
                skipped.append({"op": mapped, "reason": s.reason})
                op_detail.append({"op": mapped, "status": "skipped", "reason": s.reason})
                return None
            except Exception as e:
                LOG(f"operation {op} failed: {e}\n{traceback.format_exc()}")
                failed.append(f"OP_FAILED:{op}: {e}")
                op_detail.append({"op": mapped, "status": "failed", "error": str(e)})
                return None

        toc_status = None

        def _normalize_font():
            formatter.update_style_definitions(doc, effective)  # 双轨制轨道①
            return formatter.normalize_fonts(doc, zones, effective, cfg["options"])

        def _format_headings():
            d = headings.apply_style_definitions(doc, effective, cfg["options"], before_analysis)
            d.update(headings.apply_to_paragraphs(doc, zones, effective, cfg["options"]))
            return d

        def _page_number():
            d = layout.add_page_numbers(doc, effective, cfg["options"], before_analysis, zones)
            if d.get("warnings"):
                warnings.extend(d.pop("warnings"))
            if d.get("skipped_existing"):
                raise OpSkip("existing PAGE fields kept (page_number_skip_existing=true)")
            return d

        def _toc():
            nonlocal toc_status
            d = toc_mod.add_toc(doc, zones, effective, cfg["options"], before_analysis, accounting)
            toc_status = d["status"]
            if d.get("warnings"):
                warnings.extend(d.pop("warnings"))
            if d["status"] == "skipped":
                raise OpSkip(d.get("reason") or "toc conditions not met")
            return d

        run_op("page_setup", lambda: layout.setup_page(doc, effective))

        def _header_footer():
            src_spec = template_spec if template_spec is not None else reference_profile
            text = None
            if src_spec:
                text = ((src_spec.get("header_footer") or {}).get("header") or {}).get("text")
            if not text:
                return {"skipped": "spec has no header text"}
            d = layout.apply_template_header(doc, text)
            if d.get("headers_preserved"):
                warnings.append(
                    "header_footer: 目标已有页眉文字，按策略保留（preserve_if_present）")
            return d
        run_op("header_footer", _header_footer)

        run_op("detect_headings", lambda: {
            "headings_detected": len(zones.headings),
            "numbering_template": zones.numbering_template,
            "candidates": len(zones.candidates)})
        run_op("format_headings", _format_headings)
        run_op("normalize_font", _normalize_font)
        run_op("normalize_paragraph", lambda: formatter.normalize_paragraphs(
            doc, zones, effective, cfg["options"]))
        run_op("punctuation_normalize",
               lambda: punctuation.apply(doc, accounting))  # §10.5 门控已在 config 校验
        run_op("normalize_lists", lambda: lists.apply(doc, zones, effective))
        run_op("format_captions", lambda: formatter.format_captions(doc, zones, effective))
        run_op("format_tables", lambda: tables.format_tables(
            doc, effective, cfg["options"], before_analysis))
        run_op("format_images", lambda: images.format_images(doc, effective, cfg["options"]))
        run_op("page_number", _page_number)
        run_op("toc", _toc)
        run_op("cleanup_blank_paragraphs",
               lambda: cleanup.run(doc, zones, cfg["options"], accounting))

        sections["config"]["operations_executed"] = executed
        sections["config"]["operations_skipped"] = skipped
        sections["config"]["operations_failed"] = failed

        # ---- 16. 保存 -> 检查 -> 护栏 -> 原子落盘 ----
        tmp_out = docx_path + ".tmp"
        os.makedirs(os.path.dirname(docx_path) or ".", exist_ok=True)
        doc.save(tmp_out)

        try:
            from docx import Document as Doc2
            out_doc = Doc2(tmp_out)  # Q01：输出必须能重新打开
        except Exception as e:
            os.remove(tmp_out)
            sections["errors"] = [f"Q01: output cannot be re-opened: {e}"]
            rep = finalize("error", C.EXIT_ERROR)
            write_report(rep, report_path)
            return rep, C.EXIT_ERROR, None

        after_guard = guard.snapshot(out_doc)
        after_analysis = analyzer.analyze(out_doc)

        guard_result = guard.verify(before_guard, after_guard, accounting)
        sections["content_guard"] = {
            "enabled": cfg["content_protection"],
            "fingerprint_before": before_guard["para_fingerprint"],
            "fingerprint_after": after_guard["para_fingerprint"],
            "content_changed": guard_result["content_changed"] if cfg["content_protection"]
            else bool(accounting.get("punctuation_changes")),
            "excluded_generated": {
                "paragraphs": accounting.get("toc_paragraphs_added", 0),
                "preview": accounting.get("excluded_generated_texts", [])[:5],
            },
            "punctuation_changes": accounting.get("punctuation_changes", []),
            "integrity": guard_result["integrity"],
            "diffs": guard_result["diffs"],
        }

        checks = checker.run_checks(out_doc, {
            "before_analysis": before_analysis,
            "after_analysis": after_analysis,
            "accounting": accounting,
            "effective": effective,
            "requested": {
                "toc": "toc" in ops and toc_status == "added",
                "page_number": "page_number" in ops and not any(
                    s.get("op") == "page_number_added" for s in skipped),
                "format_headings": "format_headings" in ops,
            },
        })
        checks.insert(0, {"id": "Q01", "status": "pass",
                          "detail": "output re-opened with python-docx"})
        sections["quality_checks"] = checks

        # 护栏失败 -> 删除临时输出，不产出文件（§17.5）
        if cfg["content_protection"] and not guard_result["ok"]:
            os.remove(tmp_out)
            sections["errors"] = [f"GUARD_FAILED: {m}" for m in guard_result["integrity_failures"]] + [
                f"content diff: {d}" for d in guard_result["diffs"][:5]]
            rep = finalize("guard_failed", C.EXIT_GUARD_FAILED)
            write_report(rep, report_path)
            return rep, C.EXIT_GUARD_FAILED, None

        # 原子落盘
        os.replace(tmp_out, docx_path)
        sections["output"]["written"] = True
        sections["output"]["size_bytes"] = os.path.getsize(docx_path)

        # ---- 状态归并 ----
        error_checks = [c for c in checks if c["status"] == "fail"]
        for c in error_checks:
            warnings.append(f"{c['id']}: {c['detail']}")
        errors = failed[:] + [f"{c['id']}: {c['detail']}" for c in error_checks]

        tc = before_analysis.get("tracked_changes", {})
        if tc.get("ins") or tc.get("del"):
            warnings.append(f"文档包含未接受的修订（ins={tc['ins']}, del={tc['del']}），"
                            "建议先在 Word 中接受修订再排版")
        if before_analysis.get("comment_count"):
            warnings.append(f"文档包含 {before_analysis['comment_count']} 条批注，未处理亦未破坏")
        if before_analysis.get("cjk_char_ratio", 1.0) < 0.3:
            warnings.append("文档以西文为主，当前模板以中文排版为假设，请确认字体设置")

        sections["zones"] = zones.to_report()
        sections["analysis"] = before_analysis
        if "diagnose" in ops:
            sections["diagnosis"] = diagnose_mod.diagnose(out_doc, after_analysis, zones)
        sections["statistics"] = {
            "before": {
                "paragraphs": before_analysis["paragraph_count"],
                "tables": before_analysis["table_count"],
                "inline_images": before_analysis["inline_image_count"],
                "headings": (before_analysis["heading_counts"].get("h1", 0)
                             + before_analysis["heading_counts"].get("h2", 0)
                             + before_analysis["heading_counts"].get("h3", 0)),
                "hyperlinks": before_analysis["hyperlink_count"],
            },
            "after": {
                "paragraphs": after_analysis["paragraph_count"],
                "tables": after_analysis["table_count"],
                "inline_images": after_analysis["inline_image_count"],
                "headings": (after_analysis["heading_counts"].get("h1", 0)
                             + after_analysis["heading_counts"].get("h2", 0)
                             + after_analysis["heading_counts"].get("h3", 0)),
                "hyperlinks": after_analysis["hyperlink_count"],
            },
            "reconciliation": {
                "blank_deleted": accounting.get("blank_deleted", 0),
                "toc_added": accounting.get("toc_paragraphs_added", 0),
                "delta_explained": True,
            },
        }
        sections["operations_detail"] = op_detail
        boundary = {
            "floating_images": before_analysis["floating_image_count"],
            "textboxes": before_analysis["textbox_count"],
            "equations": before_analysis["omml_equation_count"],
            "footnotes": before_analysis["footnote_count"],
        }
        sections["untouched_boundary"] = (dict(boundary, note="以上对象未处理亦未破坏（§2.5 边界声明）")
                                          if any(boundary.values())
                                          else {"note": "未检测到边界对象"})
        sections["warnings"] = warnings
        sections["errors"] = errors


    finally:
        # 任何路径（dry-run/护栏失败/异常）都清理临时工作目录
        shutil.rmtree(workdir, ignore_errors=True)

    # ---- 渲染级验证（Phase 7 增强：真 Word 打开临时副本，更新域/验目录/数页数） ----
    if cfg.get("_render_check"):
        import render_check
        pdf_path = os.path.join(os.path.dirname(docx_path),
                                os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")             if cfg.get("_pdf") else None
        rc = render_check.render_check(docx_path, export_pdf=pdf_path)
        sections["render_check"] = rc
        if rc.get("available"):
            toc = rc.get("toc", {})
            if toc.get("tables_of_contents", 0) > 0 and not toc.get("has_real_entries"):
                warnings.append("render_check: TOC 域更新后无条目，请检查标题层级")
            if not rc.get("page_number_rendered"):
                warnings.append("render_check: 页脚未渲染出页码")
        else:
            warnings.append(f"render_check 不可用（不影响结果）: {rc.get('reason', '')[:120]}")

    status = "partial_success" if errors else "success"
    code = C.EXIT_PARTIAL if errors else C.EXIT_SUCCESS
    rep = finalize(status, code)
    write_report(rep, report_path)
    return rep, code, docx_path


# ---------------------------------------------------------------------------
# 错误报告
# ---------------------------------------------------------------------------
def error_report(code: str, message: str, input_path: str = None) -> tuple:
    from report import build_report
    sections = {
        "input": {"file": os.path.basename(input_path) if input_path else None,
                  "path": os.path.abspath(input_path) if input_path else None},
        "errors": [f"{code}: {message}"],
    }
    return build_report("error", C.EXIT_ERROR, sections), C.EXIT_ERROR, None


def try_legacy_convert(input_path: str, outdir: str) -> tuple:
    """.doc/.wps -> .docx。返回 (converted_path|None, note|error_dict)。"""
    import legacy_convert
    os.makedirs(outdir, exist_ok=True)
    base = os.path.splitext(os.path.basename(input_path))[0]
    converted = os.path.join(outdir, f"{base}_converted.docx")
    try:
        info = legacy_convert.convert_to_docx(input_path, converted)
        return converted, {"converted_from": os.path.abspath(input_path), **info}
    except legacy_convert.LegacyConvertError as e:
        return None, {"error": f"{e.reason}", "hint": e.hint}


def run_one(input_path: str, args) -> tuple:
    from report import write_report
    try:
        cfg = load_config(args.config)
        cfg["_dry_run"] = args.dry_run
        cfg["_render_check"] = getattr(args, "render_check", False)
        cfg["_pdf"] = getattr(args, "pdf", False)
        cfg.setdefault("output_dir", args.outdir)
        # 旧格式：COM 可用则先转换为 .docx（Phase 7）
        if os.path.splitext(input_path)[1].lower() in (".doc", ".wps"):
            converted, note = try_legacy_convert(input_path, args.outdir)
            if converted is None:
                rep, code, _ = error_report(
                    C.UNSUPPORTED_FILE,
                    f"legacy format conversion failed: {note.get('error')}. {note.get('hint', '')}",
                    input_path)
                return rep, code, None
            LOG(f"converted legacy format -> {converted}")
            cfg["_converted_note"] = note
            input_path = converted
        code_err, msg = validate_input(input_path)
        if code_err:
            rep, code, _ = error_report(code_err, msg, input_path)
            _, report_path = output_paths(cfg, input_path, args.outdir)
            write_report(rep, report_path)
            return rep, code, None
        return process_document(input_path, cfg)
    except ConfigError as e:
        return error_report(e.code, e.message, input_path)
    except Exception as e:
        LOG(traceback.format_exc())
        return error_report("INTERNAL_ERROR", str(e), input_path)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
SUBCOMMANDS = ("analyze", "apply", "audit", "plan", "render", "verify")


def main(argv=None):
    """入口：首参数是子命令则走子命令模式，否则按 V1 兼容模式（= apply）。"""
    argv = list(sys.argv[1:]) if argv is None else list(argv)
    if argv and argv[0] in SUBCOMMANDS:
        return _subcommand_main(argv)
    return _legacy_main(argv)


def _emit_and_save(rep, code, outdir, stem, kind):
    from report import write_report, emit_stdout
    os.makedirs(outdir, exist_ok=True)
    path = os.path.join(outdir, f"{stem}_{kind}_report.json")
    write_report(rep, path)
    emit_stdout(rep)
    return code


def _build_report_safe(status, code, sections):
    from report import build_report
    return build_report(status, code, sections)


def _run_analyze(args):
    """analyze = 只读诊断（diagnose 任务，DESIGN_V15 §13/§15）。"""
    code_err, msg = validate_input(args.target)
    if code_err:
        rep, code, _ = error_report(code_err, msg, args.target)
        return _emit_and_save(rep, code, args.outdir,
                              os.path.splitext(os.path.basename(args.target))[0], "analyze")
    cfg = load_config(args.config)
    cfg["task"] = "diagnose"
    cfg["output_dir"] = args.outdir
    rep, code, _ = process_document(args.target, cfg)
    from report import emit_stdout
    emit_stdout(rep)  # diagnose 路径不经过 legacy main 的统一输出
    return code


def _run_verify(args):
    """verify = 独立质检（无处理前状态，diff 类检查跳过）。"""
    code_err, msg = validate_input(args.target)
    if code_err:
        rep, code, _ = error_report(code_err, msg, args.target)
        return _emit_and_save(rep, code, args.outdir,
                              os.path.splitext(os.path.basename(args.target))[0], "verify")
    from docx import Document
    import analyzer
    import checker
    doc = Document(args.target)
    analysis = analyzer.analyze(doc)
    checks = checker.run_checks(doc, {
        "before_analysis": {}, "after_analysis": analysis,
        "accounting": {}, "effective": {}, "requested": {},
        "standalone": True,
    })
    checks.insert(0, {"id": "Q01", "status": "pass",
                      "detail": "file opened with python-docx"})
    fails = [c["id"] for c in checks if c["status"] == "fail"]
    sections = {
        "input": {"file": os.path.basename(args.target),
                  "path": os.path.abspath(args.target)},
        "analysis": analysis,
        "quality_checks": checks,
        "warnings": ["standalone verify: 无处理前状态，内容 diff 类检查（Q02–Q07）已跳过"],
        "errors": [f"{qid}: {[c['detail'] for c in checks if c['id'] == qid][0]}"
                   for qid in fails],
    }
    status = "partial_success" if fails else "success"
    code = C.EXIT_PARTIAL if fails else C.EXIT_SUCCESS
    rep = _build_report_safe(status, code, sections)
    return _emit_and_save(rep, code, args.outdir,
                          os.path.splitext(os.path.basename(args.target))[0], "verify")


def _run_audit(args):
    """audit = 目标 vs 模板差异清单（只读，DESIGN_V15 §13）。"""
    code_err, msg = validate_input(args.target)
    if code_err:
        rep, code, _ = error_report(code_err, msg, args.target)
        return _emit_and_save(rep, code, args.outdir,
                              os.path.splitext(os.path.basename(args.target))[0], "audit")
    terr, tmsg = validate_input(args.template)
    if terr:
        rep, code, _ = error_report(terr, f"template invalid: {tmsg}", args.target)
        return _emit_and_save(rep, code, args.outdir,
                              os.path.splitext(os.path.basename(args.target))[0], "audit")
    from docx import Document
    import audit as audit_mod
    import template_analyzer
    target_doc = Document(args.target)
    template_spec = template_analyzer.analyze_template(Document(args.template))
    result = audit_mod.run_audit(target_doc, template_spec)
    sections = {
        "input": {"file": os.path.basename(args.target),
                  "path": os.path.abspath(args.target)},
        "template": {"file": os.path.basename(args.template),
                     "path": os.path.abspath(args.template)},
        "template_spec": template_spec,
        "audit": result,
        "warnings": [], "errors": [],
    }
    rep = _build_report_safe("success", C.EXIT_SUCCESS, sections)
    return _emit_and_save(rep, C.EXIT_SUCCESS, args.outdir,
                          os.path.splitext(os.path.basename(args.target))[0], "audit")


def _run_plan(args):
    """plan = 生成排版计划与 formatting_spec.json，不执行（DESIGN_V15 §15）。"""
    cfg = load_config(args.config)
    cfg["_dry_run"] = True
    cfg.setdefault("output_dir", args.outdir)
    if getattr(args, "template", None):
        cfg["mode"] = "template"
        cfg["template"] = args.template
    elif getattr(args, "reference", None):
        cfg["mode"] = "reference"
        cfg["reference"] = args.reference
    if getattr(args, "intent", None):
        cfg["intent"] = json.loads(args.intent)
        cfg.setdefault("mode", "prompt")
        if cfg["mode"] == "auto":
            cfg["mode"] = "prompt"
    code_err, msg = validate_input(args.target)
    if code_err:
        rep, code, _ = error_report(code_err, msg, args.target)
        return _emit_and_save(rep, code, args.outdir,
                              os.path.splitext(os.path.basename(args.target))[0], "plan")
    rep, code, _ = process_document(args.target, cfg)
    from report import emit_stdout
    emit_stdout(rep)
    return code


def _run_render(args):
    """render = 渲染级验证（显式请求；COM 不可用时 exit 1）。"""
    import render_check
    pdf_path = None
    if args.pdf:
        stem = os.path.splitext(os.path.basename(args.target))[0]
        os.makedirs(args.outdir, exist_ok=True)
        pdf_path = os.path.join(args.outdir, f"{stem}.pdf")
    rc = render_check.render_check(args.target, export_pdf=pdf_path)
    sections = {
        "input": {"file": os.path.basename(args.target),
                  "path": os.path.abspath(args.target)},
        "render_check": rc,
        "warnings": [],
        "errors": [],
    }
    if rc.get("available"):
        status, code = "success", C.EXIT_SUCCESS
    else:
        status, code = "error", C.EXIT_ERROR
        sections["errors"].append(f"RENDER_UNAVAILABLE: {rc.get('reason', '')}")
    rep = _build_report_safe(status, code, sections)
    return _emit_and_save(rep, code, args.outdir,
                          os.path.splitext(os.path.basename(args.target))[0], "render")


def _subcommand_main(argv):
    ap = argparse.ArgumentParser(
        prog="docbeauty",
        description="Deterministic Word formatting engine (content-safe)")
    sub = ap.add_subparsers(dest="command", required=True)

    def _add_target_args(p):
        p.add_argument("target", help="input .docx file")
        p.add_argument("-c", "--config", help="config JSON file")
        p.add_argument("-o", "--outdir", default="./output", help="output directory")
        p.add_argument("--report", help="explicit report json path")

    p_apply = sub.add_parser("apply", help="format a document (default mode)")
    _add_target_args(p_apply)
    p_apply.add_argument("--dry-run", action="store_true")
    p_apply.add_argument("--batch", help="process every .docx in DIR")
    p_apply.add_argument("--render-check", action="store_true")
    p_apply.add_argument("--pdf", action="store_true")

    p_analyze = sub.add_parser("analyze", help="diagnose format issues (read-only)")
    _add_target_args(p_analyze)

    p_verify = sub.add_parser("verify", help="standalone quality checks on a docx")
    _add_target_args(p_verify)

    p_audit = sub.add_parser("audit", help="target vs template diff report (read-only)")
    _add_target_args(p_audit)
    p_audit.add_argument("--template", required=True, help="template .docx")

    p_plan = sub.add_parser("plan", help="build formatting plan without applying")
    _add_target_args(p_plan)
    p_plan.add_argument("--template", help="template .docx (mode=template)")
    p_plan.add_argument("--reference", help="reference .docx (mode=reference)")
    p_plan.add_argument("--intent", help="intent JSON string (mode=prompt)")

    p_render = sub.add_parser("render",
                              help="render via Word COM: update fields, verify TOC/pages")
    _add_target_args(p_render)
    p_render.add_argument("--pdf", action="store_true", help="export rendered PDF")

    args = ap.parse_args(argv)
    if args.command == "apply":
        return _legacy_main(argv[1:])  # 去掉子命令词，参数与兼容模式一致
    if args.command == "analyze":
        return _run_analyze(args)
    if args.command == "verify":
        return _run_verify(args)
    if args.command == "audit":
        return _run_audit(args)
    if args.command == "plan":
        return _run_plan(args)
    if args.command == "render":
        return _run_render(args)
    ap.error(f"unknown command: {args.command}")


def _legacy_main(argv=None):
    ap = argparse.ArgumentParser(
        prog="docbeauty",
        description="Deterministic Word formatting engine (content-safe)")
    ap.add_argument("input", nargs="?", help="input .docx file")
    ap.add_argument("-c", "--config", help="config JSON file")
    ap.add_argument("-o", "--outdir", default="./output", help="output directory")
    ap.add_argument("--report", help="explicit report json path")
    ap.add_argument("--dry-run", action="store_true", help="analyze + plan, modify nothing")
    ap.add_argument("--batch", help="process every .docx in DIR")
    ap.add_argument("--print-schema", action="store_true")
    ap.add_argument("--render-check", action="store_true",
                    help="open output in real Word (COM): update fields, verify TOC/pages")
    ap.add_argument("--pdf", action="store_true",
                    help="with --render-check: export rendered PDF (fields updated)")
    ap.add_argument("--version", action="store_true")
    args = ap.parse_args(argv)

    if args.version:
        print(C.SKILL_VERSION)
        return 0
    if args.print_schema:
        json.dump(load_schema(), sys.stdout, ensure_ascii=False, indent=2)
        return 0
    if not args.input and not args.batch:
        ap.error("input file or --batch is required")

    from report import emit_stdout

    if args.batch:
        results, worst = [], 0
        os.makedirs(args.outdir, exist_ok=True)
        for f in sorted(glob.glob(os.path.join(args.batch, "*.docx"))):
            LOG(f"batch: {f}")
            rep, code, _ = run_one(f, args)
            results.append({"file": os.path.basename(f),
                            "status": rep.get("status"), "exit_code": code})
            worst = max(worst, code)
        summary_path = os.path.join(args.outdir, "batch_summary.json")
        with open(summary_path, "w", encoding="utf-8") as fh:
            json.dump({"files": results, "worst_exit_code": worst}, fh,
                      ensure_ascii=False, indent=2)
        emit_stdout({"batch_summary": results, "worst_exit_code": worst,
                     "summary_path": summary_path})
        return worst

    rep, code, _ = run_one(args.input, args)
    emit_stdout(rep)
    return code


if __name__ == "__main__":
    sys.exit(main())
